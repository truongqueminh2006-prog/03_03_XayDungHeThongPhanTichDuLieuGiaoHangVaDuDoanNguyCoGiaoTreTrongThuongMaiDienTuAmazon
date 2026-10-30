
from pathlib import Path

from openpyxl import load_workbook
from pymongo import MongoClient, ReplaceOne

from app.core.config import settings


DATA_FILE = (
    Path(__file__).resolve().parent.parent
    / "data"
    / "Star_Schema_5_Sheets_Delivery_Data.xlsx"
)

# Sheet -> MongoDB collection -> unique key
SHEET_CONFIG = [
    ("FactOrders", "orders", "Order_ID"),
    ("DimDate", "dim_date", "DateKey"),
    ("DimVehicle", "dim_vehicle", "VehicleKey"),
    ("DimWeather", "dim_weather", "WeatherKey"),
    ("DimCategory", "dim_category", "CategoryKey"),
]


def read_sheet(workbook, sheet_name):
    worksheet = workbook[sheet_name]
    rows = worksheet.iter_rows(values_only=True)

    headers = list(next(rows))

    documents = []
    for row in rows:
        document = dict(zip(headers, row))

        # Remove empty Excel cells while preserving valid values such as 0.
        document = {
            key: value
            for key, value in document.items()
            if value is not None
        }

        documents.append(document)

    return documents


def main():
    if not DATA_FILE.exists():
        raise FileNotFoundError(
            f"Khong tim thay file du lieu: {DATA_FILE}"
        )

    workbook = load_workbook(
        DATA_FILE,
        read_only=True,
        data_only=True,
    )

    client = MongoClient(
        settings.MONGODB_URI,
        serverSelectionTimeoutMS=5000,
    )

    try:
        client.admin.command("ping")
        db = client[settings.MONGODB_DATABASE]

        print(f"Database: {settings.MONGODB_DATABASE}")

        for sheet_name, collection_name, key_field in SHEET_CONFIG:
            documents = read_sheet(workbook, sheet_name)

            # Refuse to import if a required key is missing.
            if any(
                document.get(key_field) is None
                for document in documents
            ):
                raise ValueError(
                    f"{sheet_name}: co ban ghi thieu {key_field}"
                )

            # Enforce uniqueness of each document's natural key.
            collection = db[collection_name]
            collection.create_index(key_field, unique=True)

            # Upsert by natural key: rerunning will update existing
            # records rather than insert duplicate records.
            operations = [
                ReplaceOne(
                    {key_field: document[key_field]},
                    document,
                    upsert=True,
                )
                for document in documents
            ]

            if operations:
                result = collection.bulk_write(
                    operations,
                    ordered=False,
                )
                print(
                    f"{sheet_name} -> {collection_name}: "
                    f"doc={len(documents)}, "
                    f"insert={result.upserted_count}, "
                    f"update={result.modified_count}, "
                    f"total={collection.count_documents({})}"
                )
            else:
                print(
                    f"{sheet_name}: khong co du lieu de import"
                )

        print("Import hoan tat.")

    finally:
        workbook.close()
        client.close()


if __name__ == "__main__":
    main()
